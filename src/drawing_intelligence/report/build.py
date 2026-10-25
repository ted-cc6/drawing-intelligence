"""Stage 5: the review as an engineer reads it.

An Excel checklist in the client's columns, and the sheet with every finding boxed and numbered (row n in
Excel is box n on the drawing). Code only: it reads the verified `ReviewResult` and adds nothing to it.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw

from drawing_intelligence.extraction.template import Template
from drawing_intelligence.report.zones import markers, zone
from drawing_intelligence.review.rulebook import Rulebook
from drawing_intelligence.schemas import Extraction, Finding, PageExtraction, ReviewResult
from drawing_intelligence.schemas.review import CATEGORY_LABELS
from drawing_intelligence.viz.annotate import _font, _scaled

REPORT_FILE = "review_report.xlsx"
# The client's checklist columns (data/client_set/Check List.docx), then where each row comes from.
COLUMNS = (
    "No.", "Check Category", "Requirement", "Pass/Fail", "Severity", "Comment", "Drawing Location",
    "Recommended Correction", "Rule", "Checked by",
)  # fmt: skip
WIDTHS = (6, 18, 45, 14, 10, 60, 30, 45, 10, 18)
RESULT_TEXT = {"fail": "Fail", "warning": "Warning (confirm)", "pass": "Pass"}
FILLS = {"Fail": "F8CBAD", "Warning (confirm)": "FFE699", "Pass": "C6EFCE"}
RED, ORANGE = (220, 0, 0), (230, 120, 0)
HOW_TO_READ = (
    "A Fail rests on the exact text or geometry of the PDF. A 'Warning (confirm)' is raised by an AI "
    "reviewer or asks for a judgement: an engineer should confirm it. Rules not checked are listed with the "
    "reason; none of them is counted as a pass. The rulebook is a draft until PARI confirms it."
)


@dataclass(frozen=True)
class NumberedFinding:
    no: int
    finding: Finding


def numbered(result: ReviewResult) -> list[NumberedFinding]:
    """Findings in report order (fails first, then by severity), numbered from 1."""
    return [NumberedFinding(i, f) for i, f in enumerate(result.by_severity(), 1)]


def location(f: Finding, extraction: Extraction, template: Template) -> str:
    """Where to look: sheet, zone and what is cited, e.g. "Sheet 1, zone C4: D12 '⌀ 16 H7 (+0.018/0)'"."""
    pages = {p.page_number: p for p in extraction.pages}
    parts = []
    for e in f.evidence[:3]:
        page = pages.get(e.page)
        z = zone(page, template, e.bbox) if page else None
        quote = f"'{e.quote}'" if e.quote and e.quote != e.ref else ""
        what = " ".join(x for x in (e.ref or "", quote) if x) or "marked area"
        parts.append(f"Sheet {e.page}" + (f", zone {z}" if z else "") + f": {what}")
    return "; ".join(parts)


def _source(f: Finding) -> str:
    return {"code": "code check", "lookup": "table lookup", "llm": f"AI reviewer ({f.reviewer})"}[f.source]


def write_excel(
    result: ReviewResult, rulebook: Rulebook, extraction: Extraction, template: Template, path: Path
) -> Path:
    """Three sheets: a summary, the checklist (findings, then passed rules), and what was not checked."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Checklist"
    ws.append(COLUMNS)
    for nf in numbered(result):
        f = nf.finding
        ws.append(
            [
                nf.no,
                CATEGORY_LABELS[f.category],
                f.requirement,
                RESULT_TEXT[f.result],
                f.severity.capitalize(),
                f.comment,
                location(f, extraction, template),
                f.recommended_correction or "",
                f.rule_id or "",
                _source(f),
            ]
        )
    status = {s.rule_id: s for s in result.rule_statuses}
    for rule in rulebook.rules:
        s = status.get(rule.id)
        if s and s.status == "pass":
            by = {"code": "code check", "lookup": "table lookup", "llm": "AI reviewer"}[rule.check_type]
            ws.append(
                ["", CATEGORY_LABELS[rule.category], rule.requirement, "Pass", rule.severity.capitalize(),
                 "", "", "", rule.id, by]
            )  # fmt: skip
    _style(ws, WIDTHS, result_column=4)

    nc = wb.create_sheet("Not checked")
    nc.append(("Rule", "Check Category", "Requirement", "Status", "Why"))
    for rule in rulebook.rules:
        s = status.get(rule.id)
        if s and s.status in ("not_checked", "not_applicable"):
            nc.append(
                (
                    rule.id,
                    CATEGORY_LABELS[rule.category],
                    rule.requirement,
                    s.status.replace("_", " "),
                    s.note or "",
                )
            )
    _style(nc, (10, 18, 60, 14, 60))

    sm = wb.create_sheet("Summary")
    order = ("fail", "warning", "pass", "not_checked", "not_applicable")
    counts = {k: sum(1 for s in result.rule_statuses if s.status == k) for k in order}
    first = extraction.pages[0] if extraction.pages else None
    tb = first.title_block_fields if first else {}
    n_fail = sum(f.result == "fail" for f in result.findings)
    n_warn = sum(f.result == "warning" for f in result.findings)
    rule_counts = ", ".join(f"{v} {k.replace('_', ' ')}" for k, v in counts.items())
    rows = [
        ("Drawing", result.source_file),
        ("Drawing number", tb["number"].value if "number" in tb else None),
        ("Title", tb["title"].value if "title" in tb else None),
        ("Material", tb["material"].value if "material" in tb else None),
        ("Findings", f"{n_fail} fail, {n_warn} warning"),
        ("Rules", f"{len(result.rule_statuses)}: {rule_counts}"),
        ("Reviewers run", ", ".join(result.reviewers_run) or "none"),
        ("Reviewers skipped", ", ".join(result.reviewers_skipped) or "none"),
        ("Rulebook version", result.rulebook_version),
        ("Review cost (USD)", round(result.cost_usd, 4)),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("How to read it", HOW_TO_READ),
    ]
    for r in rows:
        sm.append(r)
    _style(sm, (20, 100), header=False)
    wb.move_sheet("Summary", offset=-2)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _style(ws, widths: tuple[int, ...], result_column: int | None = None, header: bool = True) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if header:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    else:
        for row in ws.iter_rows():
            row[0].font = Font(bold=True)
    if result_column:
        for row in ws.iter_rows(min_row=2):
            cell = row[result_column - 1]
            if cell.value in FILLS:
                cell.fill = PatternFill("solid", fgColor=FILLS[str(cell.value)])


def annotate_page(
    ex_dir: Path,
    page: PageExtraction,
    findings: list[NumberedFinding],
    title_text: str = "",
    width: int = 2400,
) -> Image.Image:
    """The sheet with each finding's evidence boxed and numbered: red for a fail, orange for a warning."""
    sheet, k = _scaled(ex_dir, page, width)
    band = 70  # a header above the sheet, so the legend never covers the drawing
    img = Image.new("RGB", (sheet.width, sheet.height + band), "white")
    img.paste(sheet, (0, band))
    draw, font = ImageDraw.Draw(img), _font(22)
    n_fail = sum(nf.finding.result == "fail" for nf in findings)
    title = f"{title_text}  ·  sheet {page.page_number}  ·  {n_fail} fail, {len(findings) - n_fail} warning"
    draw.text((360, 22), title, fill=(0, 0, 0), font=_font(26))
    for nf in findings:
        color = RED if nf.finding.result == "fail" else ORANGE
        boxes = [e.bbox for e in nf.finding.evidence if e.page == page.page_number and e.bbox]
        for i, b in enumerate(boxes):
            x0, y0, x1, y1 = (v * k for v in b)
            y0, y1 = y0 + band, y1 + band
            draw.rectangle([x0 - 4, y0 - 4, x1 + 4, y1 + 4], outline=color, width=4)
            if i == 0:  # the number once per finding, at its first piece of evidence
                label = str(nf.no)
                tx, ty = x0 - 4, max(0, y0 - 34)
                tw = draw.textlength(label, font=font) + 12
                draw.rectangle([tx, ty, tx + tw, ty + 30], fill=color)
                draw.text((tx + 6, ty + 3), label, fill=(255, 255, 255), font=font)
    for i, (label, color) in enumerate((("Fail", RED), ("Warning (confirm)", ORANGE))):
        draw.rectangle([20, 10 + 28 * i, 42, 32 + 28 * i], fill=color)
        draw.text((50, 10 + 28 * i), label, fill=(0, 0, 0), font=_font(20))
    return img


def build_report(
    result: ReviewResult,
    rulebook: Rulebook,
    extraction: Extraction,
    template: Template,
    ex_dir: Path,
    out_dir: Path,
) -> list[Path]:
    """Write the Excel report and one annotated image per page that has images. Returns the files written."""
    written = [write_excel(result, rulebook, extraction, template, out_dir / REPORT_FILE)]
    nfs = numbered(result)
    for page in extraction.pages:
        if page.images is None:
            continue
        on_page = [nf for nf in nfs if any(e.page == page.page_number for e in nf.finding.evidence)]
        path = out_dir / f"page_{page.page_number}_review.png"
        annotate_page(ex_dir, page, on_page, result.source_file).save(path)
        written.append(path)
    return written


__all__ = ["REPORT_FILE", "build_report", "location", "markers", "numbered", "write_excel"]
