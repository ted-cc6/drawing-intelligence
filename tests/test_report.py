"""Stage 5 report: zones read from the border, the Excel checklist and the annotated sheet (synthetic)."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from drawing_intelligence.extraction import load_template
from drawing_intelligence.graph import run_workflow
from drawing_intelligence.report.build import COLUMNS, REPORT_FILE, location, numbered, write_excel
from drawing_intelligence.report.zones import zone
from drawing_intelligence.review import load_rulebook
from drawing_intelligence.review.rulebook import new_finding
from drawing_intelligence.schemas import (
    EvidenceRef,
    Extraction,
    PageExtraction,
    ReviewResult,
    RuleStatus,
    TextLine,
)
from tests.test_workflow import SMALL, FakeModel

TEMPLATE = load_template()
BOOK = load_rulebook()
W, H = 1191.0, 842.0


def bordered_page() -> PageExtraction:
    """An A3 page with zone markers 1-8 along the top and A-F up the left side (A at the bottom)."""
    cols = [
        TextLine(text=str(i), bbox=(100 + 141 * (i - 1), 16, 107 + 141 * (i - 1), 28)) for i in range(1, 9)
    ]
    rows = [TextLine(text=c, bbox=(28, 761 - 143 * i, 36, 773 - 143 * i)) for i, c in enumerate("ABCDEF")]
    return PageExtraction(page_number=1, page_size_pts=(W, H), rotation=0, body_text=cols + rows)


def test_zone_from_the_border_markers():
    page = bordered_page()
    assert zone(page, TEMPLATE, [380, 480, 400, 500]) == "C3"  # x near column 3, y near row C
    assert zone(page, TEMPLATE, [1080, 770, 1100, 780]) == "A8"  # bottom right, above the title block
    assert (
        zone(PageExtraction(page_number=1, page_size_pts=(W, H), rotation=0), TEMPLATE, [1, 1, 2, 2]) is None
    )


def result_with(*findings, statuses=()) -> ReviewResult:
    return ReviewResult(
        source_file="t.pdf", rulebook_version="v", findings=list(findings), rule_statuses=list(statuses)
    )


def ev(kind="dimension", **kw) -> EvidenceRef:
    base = {"kind": kind, "page": 1, "provenance": "pdf_text", "bbox": [380, 480, 400, 500]}
    return EvidenceRef(**{**base, **kw})


def test_location_names_the_sheet_zone_and_what_is_cited():
    ex = Extraction(source_file="t.pdf", template="wipro_pari", n_pages=1, pages=[bordered_page()])
    f = new_finding(BOOK.get("CN-001"), reviewer="checks", evidence=[ev(ref="D3", quote="35")], comment="c")
    assert location(f, ex, TEMPLATE) == "Sheet 1, zone C3: D3 '35'"
    t = new_finding(
        BOOK.get("TB-006"), reviewer="checks", evidence=[ev("text", quote="SCALE 1:1")], comment="c"
    )
    assert location(t, ex, TEMPLATE) == "Sheet 1, zone C3: 'SCALE 1:1'"


def test_excel_has_the_clients_columns_numbered_findings_passes_and_what_was_not_checked(tmp_path: Path):
    ex = Extraction(source_file="t.pdf", template="wipro_pari", n_pages=1, pages=[bordered_page()])
    fail = new_finding(
        BOOK.get("CN-001"), reviewer="checks", evidence=[ev(ref="D3", quote="35")], comment="off by 5"
    )
    warn = new_finding(
        BOOK.get("MC-001"), reviewer="manufacturability", evidence=[ev(ref="D7")], comment="sharp"
    )
    statuses = [
        RuleStatus(rule_id="CN-001", status="fail", finding_ids=[fail.id]),
        RuleStatus(rule_id="MC-001", status="warning", finding_ids=[warn.id]),
        RuleStatus(rule_id="TB-001", status="pass"),
        RuleStatus(rule_id="TB-003", status="not_checked", note="no title block read"),
    ]
    path = write_excel(result_with(warn, fail, statuses=statuses), BOOK, ex, TEMPLATE, tmp_path / REPORT_FILE)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Summary", "Checklist", "Not checked"]
    rows = list(wb["Checklist"].iter_rows(values_only=True))
    assert rows[0] == COLUMNS
    assert [(r[0], r[3], r[8]) for r in rows[1:3]] == [
        (1, "Fail", "CN-001"),
        (2, "Warning (confirm)", "MC-001"),
    ]
    assert rows[2][9] == "AI reviewer (manufacturability)"
    assert ("Pass", "TB-001") in {(r[3], r[8]) for r in rows[3:]}
    not_checked = list(wb["Not checked"].iter_rows(values_only=True))
    assert ("TB-003", "not checked", "no title block read") in {(r[0], r[3], r[4]) for r in not_checked}


def test_fails_come_before_warnings_whatever_the_order_found():
    fail = new_finding(
        BOOK.get("TB-006"), reviewer="checks", evidence=[ev(ref="scale", kind="title_block")], comment="c"
    )
    warn = new_finding(BOOK.get("MC-001"), reviewer="manufacturability", evidence=[ev(ref="D7")], comment="c")
    assert [nf.finding.rule_id for nf in numbered(result_with(warn, fail))] == ["TB-006", "MC-001"]


def test_review_export_writes_the_report_and_annotated_sheet(make_pdf, settings, tmp_path):
    res = run_workflow(
        make_pdf(overall="405"), settings, imaging=SMALL, completion_fn=FakeModel(), review=True,
        export_dir=tmp_path / "x",
    )  # fmt: skip
    assert res.exported_to is not None and res.review is not None
    assert (res.exported_to / REPORT_FILE).exists() and (res.exported_to / "page_1_review.png").exists()
    rows = list(
        openpyxl.load_workbook(res.exported_to / REPORT_FILE)["Checklist"].iter_rows(values_only=True)
    )
    assert sum(1 for r in rows[1:] if r[0]) == len(res.review.findings)
